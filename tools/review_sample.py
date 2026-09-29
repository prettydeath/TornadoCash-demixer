#!/usr/bin/env python3
"""Build a blinded manual-review sheet from the placebo target runs.

The sheet lists depositor / candidate address pairs in random order with
explorer links and empty verdict columns. It does not show the band, score,
signals or whether a row is a lead at all: a share of the rows are controls,
random recipients from the same search window that the tool did not flag.
The reviewer marks each row "same owner", "different owner" or "unknown"
from independent evidence (exchange deposit addresses, labels, casework).

The key (row id -> kind, band, signals) is written separately and must not
be opened before the review is finished. ``--score``
then reports precision per band with Wilson intervals, and the
"same owner" rate on controls as the reviewer's false-confirmation check.

Usage
-----
    python tools/review_sample.py --out review.xlsx
    python tools/review_sample.py --score review.xlsx
"""

from __future__ import annotations

import argparse
import json
import math
import os
import random
import sys
from collections import Counter, defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from ens_labels import CACHE  # noqa: E402

from tornado_demix.heuristics import ranked_candidates  # noqa: E402

TARGETS = os.path.join(CACHE, "placebo", "target")
KEY = os.path.join(CACHE, "review_key.json")
VERDICTS = ("same owner", "different owner", "unknown")
EXPLORER = "https://etherscan.io/address/"


def build(n_strong_mod, n_weak, n_control, seed):
    rng = random.Random(seed)
    leads = defaultdict(list)
    controls = []
    for name in sorted(os.listdir(TARGETS)):
        with open(os.path.join(TARGETS, name), encoding="utf-8") as fh:
            data = json.load(fh)
        dep = data["wallet"]
        flagged = set()
        for r in ranked_candidates(data):
            flagged.add(r["address"])
            if r["band"] == "weak":
                kind = "weak"
            elif set(r["signals"]) & {"linked", "linked_sender"}:
                kind = "linked"  # the one family the placebo test found above chance
            else:
                kind = "strong_moderate"
            leads[kind].append(
                {"depositor": dep, "candidate": r["address"], "pool": r["pool_key"],
                 "band": r["band"], "signals": sorted(r["signals"])}
            )
        for pool, res in data.get("denoms", {}).items():
            pool_rcpts = [a for a in res.get("counts", {}) if a not in flagged and a != dep]
            if pool_rcpts:
                controls.append(
                    {"depositor": dep, "candidate": rng.choice(pool_rcpts), "pool": pool,
                     "band": "control", "signals": []}
                )
    # Stratified: every linked-family lead (few), the rest of the budget from other
    # strong/moderate leads; the key keeps the stratum for per-stratum scoring.
    linked = leads["linked"][:n_strong_mod]
    other = n_strong_mod - len(linked)
    rows = (
        linked
        + rng.sample(leads["strong_moderate"], min(other, len(leads["strong_moderate"])))
        + rng.sample(leads["weak"], min(n_weak, len(leads["weak"])))
        + rng.sample(controls, min(n_control, len(controls)))
    )
    rng.shuffle(rows)
    for i, r in enumerate(rows, 1):
        r["id"] = f"R{i:03d}"
    return rows


def write_sheet(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Перевірка"
    head = ["ID", "Пул", "Депозитор", "Кандидат", "Депозитор (Etherscan)", "Кандидат (Etherscan)",
            "Висновок", "Джерело доказу", "Коментар"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E2F3")
    for r in rows:
        ws.append([r["id"], r["pool"], r["depositor"], r["candidate"],
                   EXPLORER + r["depositor"], EXPLORER + r["candidate"], "", "", ""])
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(VERDICTS), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{len(rows) + 1}")
    for col, w in zip("ABCDEFGHI", (7, 10, 44, 44, 20, 20, 17, 28, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Інструкція")
    for line in (
        "Для кожного рядка визначте, чи належать депозитор і кандидат одному власнику.",
        "Висновок: same owner / different owner / unknown.",
        "Спирайтеся лише на незалежні докази: спільна депозитна адреса біржі, мітки, KYC зі справ,",
        "явні спільні ідентифікатори. Не використовуйте результати demix і не запускайте його на цих адресах.",
        "Частина рядків — контрольні пари, яких інструмент не позначав. Їхні номери не розкриваються.",
        "Якщо доказів недостатньо — unknown. Це нормальний і чесний результат.",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 110
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")
    wb.save(path)


def wilson(k, n, z=1.96):
    if not n:
        return None
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return [round(c - h, 3), round(c + h, 3)]


def score(path):
    from openpyxl import load_workbook

    with open(KEY, encoding="utf-8") as fh:
        key = {r["id"]: r for r in json.load(fh)}
    ws = load_workbook(path)["Перевірка"]
    tally = defaultdict(Counter)
    for row in ws.iter_rows(min_row=2, values_only=True):
        rid, verdict = row[0], (row[6] or "").strip()
        if rid in key:
            k = key[rid]
            stratum = "linked" if set(k["signals"]) & {"linked", "linked_sender"} else k["band"]
            tally[stratum][verdict or "blank"] += 1
    out = {}
    for band, c in sorted(tally.items()):
        same, diff = c["same owner"], c["different owner"]
        decided = same + diff
        out[band] = {
            **dict(c),
            "precision_decided": round(same / decided, 3) if decided else None,
            "precision_decided_95ci": wilson(same, decided),
            # unknown counted as wrong (lower) and as right (upper)
            "precision_range_all": [round(same / sum(c.values()), 3),
                                    round((same + c["unknown"]) / sum(c.values()), 3)],
        }
    print(json.dumps(out, indent=1))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", default="review.xlsx")
    ap.add_argument("--strong-moderate", type=int, default=60)
    ap.add_argument("--weak", type=int, default=20)
    ap.add_argument("--controls", type=int, default=20)
    ap.add_argument("--seed", type=int, default=11)
    ap.add_argument("--score", metavar="XLSX")
    args = ap.parse_args(argv)
    if args.score:
        score(args.score)
        return
    rows = build(args.strong_moderate, args.weak, args.controls, args.seed)
    write_sheet(rows, args.out)
    with open(KEY, "w", encoding="utf-8") as fh:
        json.dump(rows, fh, indent=1)
    print(f"{len(rows)} rows -> {args.out}; key -> {KEY}")
    print(Counter(r["band"] for r in rows))


if __name__ == "__main__":
    main()
